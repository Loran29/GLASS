"""
Build results/GLASS_ablation_results.xlsx from the verified ablation files (values, no formulas), in the style of
the Stage-2 workbooks (evaluation/loanapp_p2p/scripts/make_stage2_workbook.py: fonts, header fill, borders).

Sheets: Overview (per-variant table in the scoring order, V_full minus each variant with p-values), Per cell
(dataset x goal x variant), All runs (one row per run), Stage 1 (KPI ablation per cell), Setup (datasets,
baselines, X, goals, business context with its wording history, variants).

Sources: results/stage2/ablation_per_variant.csv, ablation_per_cell.csv, ablation_summary.csv and raw/,
results/stage1/stage1_per_cell.csv, setup/ablation_config.json, setup/x_activity.json, baselines/validation.json.
Usage (repo root): python evaluation/real_life_ablation/scripts/make_ablation_workbook.py
"""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE.parent / "loanapp_p2p" / "scripts"))
import make_stage2_workbook as w1  # noqa: E402  (style_sheet, header, BOLD, TITLE)

R2, R1 = BASE / "results" / "stage2", BASE / "results" / "stage1"
OUT = BASE / "results" / "GLASS_ablation_results.xlsx"
VARIANT_TEXT = {
    "V_full": "Everything: log profile, clarification context, knowledge base, up to 4 iterations, current rules",
    "V_no_log": "No log profile or log evidence in the Stage-2 prompt",
    "V_no_context": "No clarification chat (GLASS defaults, no constraint summary)",
    "V_no_kb": "No knowledge-base retrieval or knowledge-base material, in generation and feedback",
    "V_no_refine": "One iteration only (max_iterations = 1)",
    "V_old_refine": "Earlier refinement rules: stop at first overall 'improved', best iteration by score alone",
}
WORDING = [
    ("staffing", "1 (smoke runs)", "Each resource group may be increased by up to 20 % of its current size, at least "
     "one person; working hours may be extended by up to two hours per day.",
     "Fractional headcounts (1 -> 1.2) on the one-person profiles; the merger skips them."),
    ("staffing", "2 (check run)", "Any person's role may be reinforced by one additional person with the same skills; "
     "overall, staff may grow by at most 20 %. Working hours may be extended by up to two hours per day.",
     "GLASS's compliance check read the 20 % as a per-role cap (1 -> 2 = +100 %) and rejected +1. GLASS's "
     "constraint model has per-role maxima and a budget only (no global headcount cap; known gap)."),
    ("staffing", "3 (check run)", "Each role may be reinforced by at most one additional person with the same skills. "
     "Working hours may be extended by up to two hours per day.",
     "GLASS's compliance check still judged 1 -> 2 a violation."),
    ("immutable", "1 (smoke and check runs)", "The order of the steps and the activities themselves must stay; who "
     "does what, how long steps take and how long cases wait between steps can be improved.",
     "GLASS's feasibility check treated the activities' durations as immutable."),
]


def read(path: Path) -> list[dict]:
    return list(csv.DictReader(open(path, encoding="utf-8")))


def write_rows(ws, start: int, rows: list[list]) -> int:
    for r, values in enumerate(rows, start):
        for j, v in enumerate(values, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = w1.BOX
    return start + len(rows)


def num(v, digits=3):
    try:
        return round(float(v), digits)
    except (TypeError, ValueError):
        return v if v not in (None, "") else "–"


def overview(wb, config):
    ws = wb.active
    ws.title = "Overview"
    pv = read(R2 / "ablation_per_variant.csv")
    s = config["settings"]
    ws.cell(row=1, column=1, value="GLASS real-life ablation: per-variant results").font = w1.TITLE
    ws.cell(row=2, column=1, value=f"{len(read(R2 / 'ablation_summary.csv'))} runs: 3 datasets x 3 goals x "
            f"{len(pv)} variants x repetitions; model {s['model']}; {s['seeds']} seeds; max {s['max_iterations']} "
            "iterations. Final scenario = the loop's best iteration. Scoring order: scenario produced -> valid -> "
            "executable -> merge stability > 0 -> all KPIs improved -> DHR -> NIS -> secondary.")
    labels = [k for k in pv[0] if not k.endswith((": V_full minus variant", ": p", ": pairs"))
              and k not in ("variant", "runs", "status_ok")]
    w1.header(ws, 4, ["Measure"] + [r["variant"] for r in pv])
    rows = [["Runs"] + [r["runs"] for r in pv], ["Status ok"] + [r["status_ok"] for r in pv]]
    rows += [[lbl] + [num(r[lbl]) for r in pv] for lbl in labels]
    row = write_rows(ws, 5, rows) + 1
    ws.cell(row=row, column=1, value="V_full minus each variant (rates: percentage points, Fisher's exact test; "
            "means: mean paired difference, sign test)").font = w1.BOLD
    others = [r for r in pv if r["variant"] != "V_full"]
    w1.header(ws, row + 1, ["Measure"] + [f"{r['variant']}: difference" for r in others] + [f"{r['variant']}: p" for r in others])
    rows = [[lbl] + [num(r.get(f"{lbl}: V_full minus variant")) for r in others] + [num(r.get(f"{lbl}: p"), 4) for r in others]
            for lbl in labels]
    write_rows(ws, row + 2, rows)
    w1.style_sheet(ws, [46] + [22] * (2 * len(others)), header_rows=4)


def per_cell(wb):
    ws = wb.create_sheet("Per cell")
    rows = read(R2 / "ablation_per_cell.csv")
    w1.header(ws, 1, list(rows[0]))
    write_rows(ws, 2, [[num(v) for v in r.values()] for r in rows])
    w1.style_sheet(ws, [18, 8, 14] + [13] * (len(rows[0]) - 3))


def all_runs(wb):
    ws = wb.create_sheet("All runs")
    cols = ["Configuration", "Dataset", "Goal", "Variant", "Rep", "Status", "Iterations used", "Final iteration",
            "Produced (final)", "Valid (final)", "Executable (final)", "Merge stability (final)",
            "Admissible (final)", "Violations (final)", "KPI change % (final)", "All KPIs improved (final)",
            "DHR (final)", "NIS (final)", "CTD / RED / CWD to test log (final)", "Valid (initial)",
            "KPI change % (initial)", "Loop discarded a better initial", "Loop kept the baseline",
            "direction_inconsistent rejections", "Final modifications"]
    w1.header(ws, 1, cols)
    out = []
    for path in sorted((R2 / "raw").glob("*/rep_*.json")):
        rec = json.loads(path.read_text(encoding="utf-8"))
        ev = rec.get("evaluation") or {}
        fin, ini = ev.get("final") or {}, ev.get("initial") or {}

        def kp(e):
            return "\n".join(f"{k['kpi']}: {k['pct_change']:+.1f} % (p {k['p_value']:.3g}, "
                             f"{'improved' if k['glass_improved'] else 'not improved'})"
                             if k.get("pct_change") is not None and k.get("p_value") is not None
                             else f"{k['kpi']}: {k.get('pct_change')}" for k in e.get("kpis") or []) or "–"
        its = (rec.get("loop") or {}).get("iterations") or []
        fit = its[fin["iteration"] - 1] if fin.get("iteration") and fin["iteration"] <= len(its) else {}
        mods = "\n".join(f"{m['parameter_type']} {m['direction']}: {m['target_element']}: {m['baseline_value']} -> "
                         f"{m['proposed_value']}" for m in (fit.get("patch") or {}).get("modifications", []))
        d = fin.get("distance_to_test_log") or {}
        yn = lambda v: "yes" if v else "no"
        out.append([rec["config_id"], rec["dataset"], rec["goal"], rec["variant"], rec["rep"], rec["status"],
                    (rec.get("loop") or {}).get("iterations_used"), fin.get("iteration"),
                    yn(fin.get("scenario_produced")), yn(fin.get("valid")), yn(fin.get("executable")),
                    num(fin.get("merge_stability")), yn(fin.get("admissible")),
                    "; ".join(fin.get("violations") or []) or "none", kp(fin), yn(fin.get("all_kpis_improved")),
                    num(fin.get("DHR")), num(fin.get("NIS"), 2),
                    " / ".join(f"{d[m]:.2f}" for m in ("CTD", "RED", "CWD")) if d else "–",
                    yn(ini.get("valid")), kp(ini), yn(rec.get("loop_discarded_better_initial")),
                    yn((rec.get("loop") or {}).get("no_improving_scenario")),
                    rec.get("direction_inconsistent_rejections", 0), mods or "none"])
    write_rows(ws, 2, out)
    w1.style_sheet(ws, [24, 18, 7, 13, 5, 8, 9, 9, 9, 9, 9, 9, 9, 30, 44, 9, 7, 8, 18, 9, 44, 10, 10, 10, 60])


def stage1(wb):
    ws = wb.create_sheet("Stage 1")
    rows = read(R1 / "stage1_per_cell.csv")
    w1.header(ws, 1, list(rows[0]))
    write_rows(ws, 2, [[num(v) for v in r.values()] for r in rows])
    w1.style_sheet(ws, [18, 8, 8, 24] + [14] * (len(rows[0]) - 4))


def setup(wb, config):
    ws = wb.create_sheet("Setup")
    val = json.loads((BASE / "baselines" / "validation.json").read_text(encoding="utf-8"))["datasets"]
    row = 1
    ws.cell(row=row, column=1, value="Datasets and baselines").font = w1.TITLE
    w1.header(ws, row + 1, ["Dataset", "Baseline variant", "Cases", "Start", "X (largest mean waiting, >= 10 % of cases)",
                            "X mean waiting (h)", "Validation (one GLASS run)"])

    def vtext(v: dict) -> str:
        return (f"selection CTD {v['selection_ctd_mean']}; timers kept {v['timers_preserved']}; calendar rounding "
                f"{v['calendar_rounding']['change_pct']:+.2f} % weekly hours; mean cycle time vs test log "
                f"{v['glass_vs_test_pct']:+.1f} %, vs package simulation {v['glass_vs_package_pct']:+.1f} %; events per "
                f"case {v['glass_run']['events_per_case']} (test log {v['test_log']['events_per_case']})")
    rows = [[ds, d["baseline_variant"], d["cases"], d["start"], d["x_activity"], d["x_mean_waiting_h"], vtext(val[ds])]
            for ds, d in config["datasets"].items()]
    row = write_rows(ws, row + 2, rows) + 1
    ws.cell(row=row, column=1, value="Goals and expected KPIs").font = w1.TITLE
    w1.header(ws, row + 1, ["Configuration", "Goal sentence", "Expected KPIs"])
    row = write_rows(ws, row + 2, [[c["config_id"], c["simulation_goal"], "; ".join(c["expected_kpis"])]
                                   for c in config["configurations"]]) + 1
    ws.cell(row=row, column=1, value="Business context (clarification answers, as sent)").font = w1.TITLE
    w1.header(ws, row + 1, ["Topic", "Answer"])
    bc = config["business_context"]
    row = write_rows(ws, row + 2, [[k, bc[k]] for k in ("staffing", "budget", "immutable")]) + 1
    ws.cell(row=row, column=1, value="Earlier wordings (changed after the smoke and check runs, before the batch)").font = w1.TITLE
    w1.header(ws, row + 1, ["Topic", "Version", "Wording", "Why it was changed"])
    row = write_rows(ws, row + 2, [list(w) for w in WORDING]) + 1
    ws.cell(row=row, column=1, value="Variants").font = w1.TITLE
    w1.header(ws, row + 1, ["Variant", "Definition"])
    write_rows(ws, row + 2, [[k, v] for k, v in VARIANT_TEXT.items()])
    w1.style_sheet(ws, [26, 60, 60, 60, 40, 14, 40], header_rows=0)


def main() -> None:
    config = json.loads((BASE / "setup" / "ablation_config.json").read_text(encoding="utf-8"))
    wb = Workbook()
    overview(wb, config)
    per_cell(wb)
    all_runs(wb)
    stage1(wb)
    setup(wb, config)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()

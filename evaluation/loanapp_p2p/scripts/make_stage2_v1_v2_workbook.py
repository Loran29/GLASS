"""
Build results/GLASS_Stage2_results_v1_v2.xlsx from the verified Stage-2 files (values, no formulas).

Sources: results/stage2_v1_vs_v2.csv and .md (comparison), results/stage2_results_v2/ (records and summary of
batch v2), setup/stage2_config.json. Wording follows the v1 workbook ("first proposal" / "after refinement").

Usage (repo root, environment with GLASS's dependencies and openpyxl):
  python evaluation/loanapp_p2p/scripts/make_stage2_v1_v2_workbook.py
"""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import make_stage2_workbook as w1  # noqa: E402  (names_for, readable, header, style_sheet, BOX, TITLE)
import run_stage2_batch as b  # noqa: E402

RES = HERE.parent / "results"
R2 = RES / "stage2_results_v2"
OUT = RES / "GLASS_Stage2_results_v1_v2.xlsx"

CMP_COLS = [  # (csv metric, header)
    ("initial_valid", "Valid scenario (first proposal)"),
    ("final_valid", "Valid scenario (after refinement)"),
    ("initial_all_kpis_improved", "All target KPIs improved, our rule (first proposal)"),
    ("final_all_kpis_improved", "All target KPIs improved, our rule (after refinement)"),
    ("initial_all_kpis_improved_glass_rule", "All target KPIs improved, GLASS's rule (first proposal)"),
    ("final_all_kpis_improved_glass_rule", "All target KPIs improved, GLASS's rule (after refinement)"),
    ("final_util_improved_glass", "Utilization improved, GLASS's rule: any decrease (after refinement)"),
    ("final_util_improved_batch", "Utilization improved, our rule: below 0.85 and lower (after refinement)"),
    ("final_valid_and_improved", "Valid and all KPIs improved, our rule (after refinement)"),
    ("reps_loop_discarded_better_initial", "Runs where the refinement made the result worse than the first proposal"),
    ("final_mean_recall", "Share of injected changes recovered (after refinement)"),
    ("final_mean_improvement_ratio_CTD", "Closeness to the healthy model, mean (after refinement)"),
    ("final_median_improvement_ratio_CTD", "Closeness to the healthy model, median (after refinement)"),
]
RUN_COLS = ["Configuration", "Process", "Problems", "Rep", "Refinement iterations used", "Final iteration",
            "Valid (first proposal)", "Valid (after refinement)", "Violations (first proposal)",
            "Violations (after refinement)", "KPI change % (first proposal)", "KPI change % (after refinement)",
            "All KPIs improved (first proposal)", "All KPIs improved (after refinement)",
            "Problems addressed (first proposal)", "Problems addressed (after refinement)",
            "Injected changes recovered (first proposal)", "Injected changes recovered (after refinement)",
            "Closeness to healthy (first proposal)", "Closeness to healthy (after refinement)",
            "Refinement made it worse", "Changes in the scenario after refinement"]


def fmt(v, nd):
    if v in ("", None):
        return "–"
    x = float(v)
    return f"{x:.{nd}f}" if nd else f"{int(round(x))}"


def put(ws, row, values):
    for j, v in enumerate(values, 1):
        ws.cell(row=row, column=j, value=v).border = w1.BOX


def main() -> None:
    b.quiet()
    config = json.loads(b.CONFIG_PATH.read_text(encoding="utf-8"))
    cmp_rows = list(csv.DictReader(open(RES / "stage2_v1_vs_v2.csv", encoding="utf-8")))
    md = (RES / "stage2_v1_vs_v2.md").read_text(encoding="utf-8")
    wb = Workbook()

    # ---------------------------------------------------------------- sheet 1: comparison
    ws = wb.active
    ws.title = "v1 vs v2 per configuration"
    w1.header(ws, 1, ["Configuration", "Has a utilization KPI"] + [h for _, h in CMP_COLS])
    r = 2
    for row in cmp_rows:
        util = row["has_utilization_kpi"]
        cells = [row["config_id"], {"True": "yes", "False": "no"}.get(util, "")]
        for m, _ in CMP_COLS:
            nd = 2 if ("mean" in m or "median" in m) else 0
            if m.startswith("final_util_") and util == "False":
                cells.append("–")
            else:
                cells.append(f"{fmt(row[f'{m}_v1'], nd)} / {fmt(row[f'{m}_v2'], nd)}")
        put(ws, r, cells)
        if util == "":
            for j in range(1, len(cells) + 1):
                ws.cell(row=r, column=j).font = w1.BOLD
        r += 1
    ws.cell(row=r, column=1, value="Cells: v1 / v2. Counts are out of 5 runs per configuration (70 per batch in the "
            "total rows). Recall and closeness are means over the evaluated runs, closeness also as a median. "
            "Source: stage2_v1_vs_v2.csv (from the verified per-configuration and summary CSVs).")
    r += 2
    ws.cell(row=r, column=1, value="Change from v1 to v2 per group (v2 − v1)").font = w1.TITLE
    w1.header(ws, r + 1, ["Group", "Valid and all KPIs improved, our rule (after refinement)",
                          "Runs where the refinement made the result worse than the first proposal",
                          "Share of injected changes recovered (after refinement)",
                          "Closeness to the healthy model, mean (after refinement)",
                          "Closeness to the healthy model, median (after refinement)"])
    groups = [x for x in cmp_rows if x["has_utilization_kpi"] == "" and not x["config_id"].startswith("TOTAL")]
    for k, g in enumerate(groups, r + 2):
        label = "With a utilization KPI" if "with a" in g["config_id"] else "Without a utilization KPI (control)"
        put(ws, k, [label, f"{float(g['final_valid_and_improved_delta']):+g}",
                    f"{float(g['reps_loop_discarded_better_initial_delta']):+g}",
                    f"{float(g['final_mean_recall_delta']):+.3f}",
                    f"{float(g['final_mean_improvement_ratio_CTD_delta']):+.3f}",
                    f"{float(g['final_median_improvement_ratio_CTD_delta']):+.3f}"])
    r += 2 + len(groups) + 1
    ws.cell(row=r, column=1, value="Control check").font = w1.TITLE
    check = md.split("**Control check**", 1)[1].split("\n")[1:]
    for line in [l.strip().lstrip("- ").replace("**", "") for l in check if l.strip().startswith("- ")]:
        r += 1
        ws.cell(row=r, column=1, value=line)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    w1.style_sheet(ws, [22, 11] + [15] * len(CMP_COLS))

    # ---------------------------------------------------------------- sheet 2: all 70 v2 runs
    ws2 = wb.create_sheet("All 70 runs (v2)")
    w1.header(ws2, 1, RUN_COLS)
    summ = list(csv.DictReader(open(R2 / "stage2_summary.csv", encoding="utf-8")))
    cfgs = {c["config_id"]: c for c in config["configurations"]}
    order = sorted(cfgs, key=lambda c: (c.split("_")[0], len(c.split("_")[1]), c))
    names = {cid: w1.names_for(cfgs[cid], config) for cid in order}
    row = 2
    for cid in order:
        gs = cfgs[cid]["injected_problems"]
        for rep in range(1, 6):
            rec = json.loads((R2 / "raw" / cid / f"rep_{rep}.json").read_text(encoding="utf-8"))
            ev = rec.get("evaluation") or {}
            ph = {p: ev.get(p) or {"error": "no scenario"} for p in ("initial", "final")}

            def yn(e):
                return "no scenario" if "error" in e else ("yes" if e.get("valid") else "no")

            def allimp(e):
                if "error" in e:
                    return "no scenario"
                return "yes" if e.get("kpis") and all(k.get("improved") is True for k in e["kpis"]) else "no"

            def kpis(e):
                if "error" in e:
                    return "no scenario"
                return "\n".join(f"{k['kpi']}: {k['pct_change']:+.1f} % ({'improved' if k.get('improved') else 'not improved'})"
                                 if k.get("pct_change") is not None else f"{k['kpi']}: n/a" for k in e.get("kpis") or [])

            def addressed(e):
                if "error" in e:
                    return "no scenario"
                sc = e.get("scores") or {}
                return ", ".join(f"{g} {'yes' if sc.get(f'addressed_{g}') else 'no'}" for g in gs)

            def num(e, f):
                v = f(e) if "error" not in e else None
                return round(v, 4) if isinstance(v, (int, float)) else "–"

            it_no = ph["final"].get("iteration")
            fin_it = rec["loop"]["iterations"][it_no - 1] if it_no else {}
            mods = "\n".join(
                f"{m['parameter_type']} {m['direction']}: {w1.readable(m['target_element'], names[cid])}: "
                f"{w1.readable(str(m['baseline_value']), names[cid])} → {w1.readable(str(m['proposed_value']), names[cid])}"
                for m in ((fin_it or {}).get("patch") or {}).get("modifications", [])) or "–"
            srow = next(x for x in summ if x["config_id"] == cid and x["rep"] == str(rep) and x["phase"] == "final")
            put(ws2, row, [cid, rec["process"], "+".join(gs), rep, len(rec["loop"]["iterations"]), it_no or "–",
                           yn(ph["initial"]), yn(ph["final"]),
                           "–" if "error" in ph["initial"] else ("; ".join(ph["initial"].get("violations") or []) or "none"),
                           "–" if "error" in ph["final"] else ("; ".join(ph["final"].get("violations") or []) or "none"),
                           kpis(ph["initial"]), kpis(ph["final"]), allimp(ph["initial"]), allimp(ph["final"]),
                           addressed(ph["initial"]), addressed(ph["final"]),
                           num(ph["initial"], lambda e: (e.get("scores") or {}).get("recall")),
                           num(ph["final"], lambda e: (e.get("scores") or {}).get("recall")),
                           num(ph["initial"], lambda e: (e.get("improvement_ratio") or {}).get("CTD")),
                           num(ph["final"], lambda e: (e.get("improvement_ratio") or {}).get("CTD")),
                           "yes" if srow.get("loop_discarded_better_initial") == "True" else "no", mods])
            row += 1
    w1.style_sheet(ws2, [18, 12, 10, 5, 10, 9, 10, 10, 26, 26, 34, 34, 11, 11, 16, 16, 10, 10, 10, 10, 11, 60])

    # ---------------------------------------------------------------- sheet 3: settings
    ws3 = wb.create_sheet("Settings and scoring")
    w1.header(ws3, 1, ["Item", "Value"])
    s = config["settings"]
    rules = b.__doc__.split("SCORING", 1)[1].split("END SCORING")[0].strip("\n")
    sha = lambda p: b.s1.sha256(p.read_bytes())
    items = [
        ("How to read", "\"First proposal\" = the scenario GLASS proposes first; \"after refinement\" = the scenario "
         "GLASS selects after its simulation-based refinement (up to 4 iterations). On the first sheet each cell is "
         "v1 / v2; counts are out of 5 runs per configuration."),
        ("Batches", "v1 and v2 differ only in the GLASS version. Both use the same batch script, configuration, "
         "seeds and baseline logic; baselines and healthy logs were recomputed for v2."),
        ("GLASS commit, batch v1", "8039058"),
        ("GLASS commit, batch v2", "5112fbc"),
        ("Batch script (both batches)", f"scripts/run_stage2_batch.py, SHA-256 {sha(HERE / 'run_stage2_batch.py')}"),
        ("Configuration (both batches)", f"setup/stage2_config.json, SHA-256 {sha(b.CONFIG_PATH)}"),
        ("Metric change from v1 to v2 (GLASS Resource Utilization)",
         "v1: per resource, busy time ÷ simulated span (first start to last end, wall-clock time); the KPI is the "
         "mean over all resources. When a fix clears a backlog the span shrinks and this value rises, so GLASS's "
         "loop read correct fixes as degradations. v2: per resource, busy time within working hours ÷ available "
         "working time from its calendar; the KPI is evaluated for the role it names (e.g. Loan Officer), with the "
         "all-resource mean kept as a detail."),
        ("Model", s["model"]),
        ("Temperature schedule", "0.3 for the initial proposal; GLASS raises it by 0.1 per iteration "
         "(min(0.3 + 0.1 × (iteration − 1), 0.7)): 0.3 / 0.4 / 0.5 / 0.6. GLASS's context-summary and compliance "
         "calls run at 0.1."),
        ("Maximum iterations", s["max_iterations"]), ("Repetitions per configuration", s["repetitions"]),
        ("Seeds", f"{s['seeds']} (42–51), paired: baseline, first proposal and after refinement use the same seeds"),
        ("Cases per simulation", s["cases"]), ("Simulation start", s["start"]),
        ("KPI validation (lead author's interpretation)",
         "Utilization KPIs were passed with target_direction minimize (Stage 1: maximize, maintain in one case) and "
         "the sentence \"" + b.OVERLOAD_TEXT + "\" appended to the description. Rationale: in the G3 configurations "
         "the target pool is overloaded; maximizing its utilization would reward overloading it and score the "
         "recovery (more staff) as a degradation."),
        ("Utilization verdicts", "Our rule (scored): the target pool's utilization of its working time is below "
         "0.85 and lower than the baseline. GLASS's rule: GLASS's own verdict, any decrease of its utilization value "
         "(v1: all-resource mean of busy time ÷ span; v2: role-resolved working-time utilization). \"All target KPIs "
         "improved, GLASS's rule\" applies GLASS's verdict to every KPI."),
        ("Utilization KPI without numeric target",
         "The validated utilization KPI tells GLASS to minimize utilization but carries no target value. GLASS's rule "
         "counts any decrease as improved, and its loop stops at the first iteration in which all target KPIs move "
         "the right way. The 0.85 limit is an evaluation rule GLASS never received. A numeric target in the KPI text "
         "is a protocol variant for a possible later batch; it is not applied here."),
        ("Scoring rules (primary and secondary)", rules),
        ("Midnight mapping (distance measures only)",
         "Prosimos logs a task that ends exactly at the end of a working interval (e.g. 17:00) at 00:00 the next day. "
         "For the distance measures only, every log (baseline, first proposal, after refinement, healthy) maps an end "
         "of exactly 00:00:00 that directly follows the resource's working-day end on the previous day back to that "
         "end; counts are recorded per log and the raw logs are unchanged. CWD on the raw logs is also recorded."),
        ("Control check", "Configurations without a utilization KPI are the control, since the GLASS change does "
         "not target them. Outcome: valid and all KPIs improved (after refinement). A two-sided sign test runs over "
         "the configurations (ties dropped), and a pooled run-level Fisher exact test over all their runs."),
    ]
    for i, (k, v) in enumerate(items, 2):
        put(ws3, i, [k, v])
    w1.style_sheet(ws3, [38, 120])

    wb.save(OUT)
    print(f"written {OUT.relative_to(HERE.parents[2]).as_posix()}")


if __name__ == "__main__":
    main()

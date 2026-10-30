"""Compare the synthetic BPS model variants with their originals and write an Excel report.

Inputs (relative to this experiment folder)
  inputs/bps_models/<Dataset>_original.bpmn      baseline (LoanApp, Procure2Pay)
  inputs/bps_models/<Dataset>_<variant>.bpmn     resource_cont, durations, extraneous, calendars
  inputs/quality_measures.csv                    table for the 'Quality Measures' sheet

Output
  results/synthetic_vs_original_comparison.xlsx

Run from anywhere:  python scripts/build_workbook.py
"""
import csv
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from bpmn_parser import NS, parse

HERE = Path(__file__).resolve().parent
DATA = HERE.parent / 'inputs' / 'bps_models'
QM_CSV = HERE.parent / 'inputs' / 'quality_measures.csv'
OUT = HERE.parent / 'results' / 'synthetic_vs_original_comparison.xlsx'

DATASETS = ('LoanApp', 'Procure2Pay')

# Gateways in Procure2Pay with no activity directly before them: short hand-checked labels.
GATEWAY_LABELS = {'sid-112BF415-63F3-473C-B4DA-440EBF13D7D8': 'Start (first decision)',
                  'Gateway_1jc2ltt': 'Quotation decision',
                  'Gateway_1dst547': 'Before invoice release'}


def bpmn(ds, variant):
    return DATA / f'{ds}_{variant}.bpmn'


def process_graph(path):
    """Element kinds, names and sequence flows of the BPMN process."""
    p = ET.parse(path).getroot().find('b:process', NS)
    kinds = {e.get('id'): e.tag.split('}')[1] for e in p if e.get('id') and not e.tag.endswith('sequenceFlow')}
    names = {e.get('id'): (e.get('name') or '').strip() for e in p if e.get('id')}
    flows = [(s.get('sourceRef'), s.get('targetRef')) for s in p.findall('b:sequenceFlow', NS)]
    return kinds, names, flows


# ---------------------------------------------------------------- comparisons

def resource_rows():
    return [[ds, k, n, parse(bpmn(ds, 'resource_cont'))['resources'][k], None, None]
            for ds in DATASETS for k, n in parse(bpmn(ds, 'original'))['resources'].items()]


def duration_rows():
    rows = []
    for ds in DATASETS:
        o, m = parse(bpmn(ds, 'original'))['elems'], parse(bpmn(ds, 'durations'))['elems']
        for i, x in o.items():
            rows.append([ds, x['name'], x['dur']['eff_mean_min'], m[i]['dur']['eff_mean_min'], None])
    return rows


def timer_rows():
    """One row per timer in the _extraneous files: between which elements it sits and how long it delays.
    The originals have no timers: there the two elements are connected directly."""
    def fmt(d):
        mins = d['eff_mean_min']
        t = f"{mins / 60:g} h" if mins >= 60 else f"{mins:g} min"
        return t if d['type'] == 'FIXED' else f"~{t} ({d['type'].lower()})"

    rows = []
    for ds in DATASETS:
        kinds, names, fl = process_graph(bpmn(ds, 'extraneous'))

        def name(i):
            if kinds[i] == 'task':
                return names[i]
            if i in GATEWAY_LABELS:
                return GATEWAY_LABELS[i]
            pre = [names[s] for s, t in fl if t == i and kinds.get(s) == 'task']
            g = 'Parallel join' if kinds[i] == 'parallelGateway' else 'Decision'
            return f"{g} after '{' / '.join(pre)}'" if pre else f"{g} gateway"

        elems = parse(bpmn(ds, 'extraneous'))['elems']
        for eid, k in kinds.items():
            if k != 'intermediateCatchEvent':
                continue
            src = [s for s, t in fl if t == eid][0]
            tgt = [t for s, t in fl if s == eid][0]
            rows.append([ds, name(src), name(tgt), fmt(elems[eid]['dur'])])
    return rows


def calendar_rows():
    """'Business Hours' timetable in the original and _calendars files (the only one any resource uses)."""
    rows = []
    for ds in DATASETS:
        for v in ('original', 'calendars'):
            si = ET.parse(bpmn(ds, v)).getroot().find('.//q:processSimulationInfo', NS)
            t = [t for t in si.iter('{%s}timetable' % NS['q']) if t.get('name') == 'Business Hours'][0]
            x = next(t.iter('{%s}rule' % NS['q']))
            users = [r.get('name') for r in si.iter('{%s}resource' % NS['q']) if r.get('timetableId') == t.get('id')]
            rows.append([ds, f'{ds}_{v}.bpmn', f"{x.get('fromWeekDay').title()}–{x.get('toWeekDay').title()}",
                         x.get('fromTime')[:5], x.get('toTime')[:5], ', '.join(users)])
    return rows


# ---------------------------------------------------------------- workbook

FONT = 'Arial'
HEADER_FONT = Font(name=FONT, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='1F4E78')
_thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def table(ws, header, data, widths, autofilter=True):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.border = HEADER_FONT, HEADER_FILL, BORDER
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in data:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.border = Font(name=FONT), BORDER
            c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    if autofilter:
        ws.auto_filter.ref = ws.dimensions


def build():
    wb = Workbook()

    ws = wb.active
    ws.title = 'Resources'
    table(ws, ['Dataset', 'Resource', 'Original count', 'resource_cont count', 'Change', '% change'],
          resource_rows(), [13, 22, 15, 18, 10, 10])
    for i in range(2, ws.max_row + 1):
        ws.cell(i, 5).value = f'=D{i}-C{i}'
        ws.cell(i, 6).value = f'=IF(C{i}=0,0,E{i}/C{i})'
        ws.cell(i, 6).number_format = '0%;-0%;0%'

    ws = wb.create_sheet('Durations')
    table(ws, ['Dataset', 'Activity', 'Original mean (min)', 'durations mean (min)', 'Change (min)'],
          duration_rows(), [13, 42, 14, 14, 12])
    for i in range(2, ws.max_row + 1):
        ws.cell(i, 5).value = f'=D{i}-C{i}'
        ws.cell(i, 5).number_format = '0.0;[Red]-0.0;0'

    table(wb.create_sheet('Extraneous Timers'), ['Dataset', 'From', 'To', 'Delay'],
          timer_rows(), [13, 55, 40, 24], autofilter=False)

    table(wb.create_sheet('Calendars'), ['Dataset', 'File', 'Weekdays', 'From', 'To', 'Resources using it'],
          calendar_rows(), [13, 38, 18, 8, 8, 60])

    with open(QM_CSV, encoding='utf-8-sig', newline='') as f:
        qm = list(csv.reader(f))
    ws = wb.create_sheet('Quality Measures')
    table(ws, qm[0], qm[1:], [40, 50, 55, 70])
    for i in range(2, ws.max_row + 1):
        ws.cell(i, 1).font = Font(name=FONT, bold=True)

    wb.calculation.fullCalcOnLoad = True  # formulas have no cached values until Excel opens the file
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build()

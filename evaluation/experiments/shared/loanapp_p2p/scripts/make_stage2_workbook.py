"""
Build results/GLASS_Stage2_results.xlsx from the verified Stage-2 result files (generated, not retyped).

Sources: results/stage2_results/stage2_per_config.csv, stage2_summary.csv, the 70 run records (raw/), README.md;
setup/stage2_config.json; display names of activities, timers and roles from each baseline as GLASS imports it.
Totals on sheet 1 are COUNTIFS formulas over sheet 2 (Excel recalculates them when the file is opened).

Usage (repo root, environment with GLASS's dependencies and openpyxl):
  python evaluation/experiments/shared/loanapp_p2p/scripts/make_stage2_workbook.py
"""
from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
BASE = HERE.parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent))
import lp_paths as P  # noqa: E402  (experiment paths; puts the other experiments' scripts on sys.path)
R = P.OLDER / "stage2_results"
OUT = P.OLDER / "GLASS_Stage2_results.xlsx"
import run_stage2_batch as b  # noqa: E402

FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
TITLE = Font(name="Arial", size=12, bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GLASS_FIXES = [
    ("7c93a1c", "Timer / intermediate-event delays imported from the baseline, simulated and patchable (event_duration)"),
    ("893db82", "Normal-distribution std no longer treated as variance in scenario conversion"),
    ("cb9cb17", "LLM-proposed normal std no longer treated as variance"),
    ("5b3f717", "Arrival and resource calendars exported as in the baseline instead of being replaced by 24/7"),
    ("bb6cd17", "Clarification-context summary reads staffing and overtime flexibility correctly"),
    ("8039058", "Role calendar import/patching applies to every resource of a role"),
]
FINDINGS = [
    ("GLASS produces valid, KPI-improving scenarios for simple problems, but recovers only part of the injected causes.",
     "Validity is high, and single-problem LoanApp configurations improve their KPI in nearly every run. Recall is "
     "mostly 0.06–0.25; the exception is LoanApp G3 (staffing), with recall 1.0 and CTD ratio 0.87. Typical pattern: "
     "one obvious activity shortened or one timer halved, plus extra staff, often in roles that were not affected, "
     "instead of reversing all injected changes."),
    ("GLASS's utilization metric makes the refinement loop discard correct fixes in congested configurations.",
     "GLASS computes utilization as busy time ÷ the whole simulated span, so it rises when a fix clears a backlog. "
     "Under minimize, the feedback calls the correct fix a degradation. In LoanApp G1+G3 and G1+G2+G3, the initial "
     "patch (adding Loan Officers) reaches CTD ratio 0.96 / 0.88; the loop then picks near-no-op finals (CTD ratio "
     "about 0.0) in all 10 runs. Overall the loop discarded a better initial in 22 runs; in 15 of them GLASS's "
     "utilization rose while the pool's real utilization fell."),
    ("The refinement loop helps in some configurations.",
     "All KPIs improved, initial → final: LoanApp G2 3 → 4; Procure2Pay G1 1 → 3, G3 0 → 2, G1+G2 1 → 3."),
    ("Procure2Pay is harder than LoanApp.",
     "In 5 runs iteration 1 failed GLASS's validator because the LLM used the resource-id stem QBP_DEFAULT_RESOURCE "
     "as a role name; later iterations succeeded and those finals were evaluated in the completion step. Two "
     "Procure2Pay G2 finals break the 35 % budget (+39 %): GLASS shows the limit to the LLM but only enforces "
     "absolute monthly budgets. No Procure2Pay configuration combining G1 and G3 is fully improved."),
    ("Clarification context: the summariser fix holds in most cases.",
     "GLASS re-summarises the chat, including the injected feedback, in every iteration, so the constraints block "
     "differs from iteration 1 in 86 of 98 later iterations, mostly in wording. In 4 of 168 summaries it still read "
     "\"+3 per role\" as a global cap of 3."),
    ("The utilization limit is a sensitive choice.",
     "At 0.90 instead of 0.85, 22 instead of 18 initial scenarios count as improved; finals stay at 16."),
]


def style_sheet(ws, widths: list[int], header_rows: int = 1) -> None:
    for row in ws.iter_rows():
        for c in row:
            c.font = BOLD if c.font.bold else (TITLE if c.font.size == 12 else FONT)
            c.alignment = WRAP
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=header_rows + 1, column=1)


def header(ws, row: int, values: list[str]) -> None:
    for j, v in enumerate(values, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.font, c.fill, c.border = BOLD, HEAD_FILL, BOX


def names_for(cfg: dict, config: dict) -> dict[str, str]:
    """id -> display name for activities, timers, roles, resources and calendars of a configuration's baseline."""
    inputs = b.load_inputs(cfg, config)
    scn = b.load_scenario(inputs["sim_params_text"], inputs["bpmn"]).model_dump(mode="json")
    names = {}
    for m in scn["models"]:
        names.update({a["id"]: a["name"] for a in m["modelParameter"]["activities"] if a.get("name")})
        names.update({e["id"]: e["name"] for e in m["modelParameter"].get("intermediateEvents", []) if e.get("name")})
    names.update({p["id"]: p["name"] for p in inputs["sim_params"]["resource_profiles"]})   # pool id -> role
    cal_roles: dict[str, list[str]] = {}
    for role in scn["resourceParameters"]["roles"]:
        cal_roles.setdefault(role["schedule"], []).append(role["id"])
        for n, res in enumerate(role["resources"], 1):
            names[res["id"]] = f"{role['id']} resource {n}"
            if res.get("calendar"):
                cal_roles.setdefault(res["calendar"], []).append(role["id"])
    for cal, roles in cal_roles.items():
        names[cal] = "calendar of " + ", ".join(dict.fromkeys(roles))
    return names


def readable(text: str, names: dict[str, str]) -> str:
    for i in sorted(names, key=len, reverse=True):
        if i and i in text and names[i] != i:
            text = re.sub(rf"(?<![\w-]){re.escape(i)}(?![\w-])", names[i], text)
    return text


def kpi_type(k: dict) -> str:
    if k.get("category") == "utilization":
        return "Utilization"
    return "Waiting time" if "wait" in k["kpi"].lower() else "Cycle time"


def main() -> None:
    b.quiet()
    config = json.loads(b.CONFIG_PATH.read_text(encoding="utf-8"))
    per = {r["config_id"]: r for r in csv.DictReader(open(R / "stage2_per_config.csv", encoding="utf-8"))}
    summ = list(csv.DictReader(open(R / "stage2_summary.csv", encoding="utf-8")))
    recs = {}
    for p in sorted((R / "raw").glob("*/rep_*.json")):
        r = json.loads(p.read_text(encoding="utf-8"))
        recs[(r["config_id"], r["rep"])] = r
    order = sorted(per, key=lambda c: (c.split("_")[0], len(c.split("_")[1]), c))
    cfgs = {c["config_id"]: c for c in config["configurations"]}
    names = {cid: names_for(cfgs[cid], config) for cid in order}

    wb = Workbook()
    # ------------------------------------------------------------- sheet 2 first (sheet 1 totals refer to it)
    ws2 = wb.active
    ws2.title = "All 70 runs"
    cols2 = ["Configuration", "Process", "Problems", "Rep", "Iterations used", "Final iteration",
             "Valid (initial)", "Valid (final)", "Violations (initial)", "Violations (final)",
             "KPI change % (initial)", "KPI change % (final)",
             "All KPIs improved (initial)", "All KPIs improved (final)",
             "Problems addressed (initial)", "Problems addressed (final)",
             "Recall (initial)", "Recall (final)", "CTD ratio (initial)", "CTD ratio (final)",
             "Discarded better initial", "Final modifications"]
    header(ws2, 1, cols2)
    type_counts = {t: {"initial": [0, 0], "final": [0, 0]} for t in ("Cycle time", "Waiting time", "Utilization")}
    row = 2
    for cid in order:
        cfg = cfgs[cid]
        gs = cfg["injected_problems"]
        for rep in range(1, 6):
            rec = recs[(cid, rep)]
            ev = rec["evaluation"]
            ph = {p: ev.get(p) or {} for p in ("initial", "final")}

            def yn(e, key):
                if "error" in e:
                    return "no scenario"
                return "yes" if e.get(key) else "no"

            def allimp(e):
                if "error" in e:
                    return "no scenario"
                return "yes" if e.get("kpis") and all(k.get("improved") is True for k in e["kpis"]) else "no"

            def kpis(e):
                if "error" in e:
                    return "no scenario"
                return "\n".join(f"{k['kpi']}: {k['pct_change']:+.1f} % ({'improved' if k.get('improved') else 'not improved'})"
                                 if k.get("pct_change") is not None else f"{k['kpi']}: n/a" for k in e.get("kpis") or [])

            def addressed(e):
                if "error" in e:
                    return "no scenario"
                sc = e.get("scores") or {}
                return ", ".join(f"{g} {'yes' if sc.get(f'addressed_{g}') else 'no'}" for g in gs)

            def num(e, f):
                v = f(e) if "error" not in e else None
                return round(v, 4) if isinstance(v, (int, float)) else "–"

            fin_it = rec["loop"]["iterations"][ph["final"]["iteration"] - 1]
            mods = "\n".join(
                f"{m['parameter_type']} {m['direction']}: {readable(m['target_element'], names[cid])}: "
                f"{readable(str(m['baseline_value']), names[cid])} → {readable(str(m['proposed_value']), names[cid])}"
                for m in (fin_it.get("patch") or {}).get("modifications", []))
            srow = next(x for x in summ if x["config_id"] == cid and x["rep"] == str(rep) and x["phase"] == "final")
            values = [cid, rec["process"], "+".join(gs), rep, len(rec["loop"]["iterations"]), ph["final"].get("iteration"),
                      yn(ph["initial"], "valid"), yn(ph["final"], "valid"),
                      "–" if "error" in ph["initial"] else ("; ".join(ph["initial"].get("violations") or []) or "none"),
                      "; ".join(ph["final"].get("violations") or []) or "none",
                      kpis(ph["initial"]), kpis(ph["final"]), allimp(ph["initial"]), allimp(ph["final"]),
                      addressed(ph["initial"]), addressed(ph["final"]),
                      num(ph["initial"], lambda e: (e.get("scores") or {}).get("recall")),
                      num(ph["final"], lambda e: (e.get("scores") or {}).get("recall")),
                      num(ph["initial"], lambda e: (e.get("improvement_ratio") or {}).get("CTD")),
                      num(ph["final"], lambda e: (e.get("improvement_ratio") or {}).get("CTD")),
                      "yes" if srow["loop_discarded_better_initial"] == "True" else "no", mods]
            for j, v in enumerate(values, 1):
                ws2.cell(row=row, column=j, value=v).border = BOX
            for p in ("initial", "final"):
                for k in (ph[p].get("kpis") or []):
                    t = type_counts[kpi_type(k)][p]
                    t[1] += 1
                    t[0] += k.get("improved") is True
            row += 1
    last = row - 1
    style_sheet(ws2, [18, 12, 10, 5, 10, 9, 10, 10, 26, 26, 34, 34, 11, 11, 16, 16, 9, 9, 9, 9, 11, 60])

    # ------------------------------------------------------------- sheet 1
    ws1 = wb.create_sheet("Results per configuration", 0)
    cols1 = ["Configuration", "Problems", "Valid (initial → final)", "All KPIs improved (initial → final)",
             "Valid and improved (final)", "Problems addressed (final)", "Precision (final)",
             "Recall (initial → final)", "CTD ratio (initial → final)", "Loop discarded a better initial",
             "Mean iterations"]
    header(ws1, 1, cols1)
    fmt = lambda x: f"{float(x):.2f}" if x not in ("", None) else "–"
    for i, cid in enumerate(order, 2):
        r = per[cid]
        gs = r["injected_problems"].split("|")
        vals = [cid, "+".join(gs), f"{r['initial_valid']} → {r['final_valid']}",
                f"{r['initial_all_kpis_improved']} → {r['final_all_kpis_improved']}", int(r["final_valid_and_improved"]),
                ", ".join(f"{g} {r.get(f'final_addressed_{g}', '')}" for g in gs), fmt(r["final_mean_precision"]),
                f"{fmt(r['initial_mean_recall'])} → {fmt(r['final_mean_recall'])}",
                f"{fmt(r['initial_mean_improvement_ratio_CTD'])} → {fmt(r['final_mean_improvement_ratio_CTD'])}",
                int(r["reps_loop_discarded_better_initial"]), float(r["mean_iterations_used"])]
        for j, v in enumerate(vals, 1):
            ws1.cell(row=i, column=j, value=v).border = BOX
    n = len(order) + 1
    ws1.cell(row=n + 1, column=1, value="Counts out of 5 runs per configuration. Precision, recall and CTD ratio are means "
             "over the evaluated runs. Source: stage2_per_config.csv (verified).")
    t0 = n + 3
    ws1.cell(row=t0, column=1, value="Totals over 70 runs").font = TITLE
    header(ws1, t0 + 1, ["Measure", "Initial", "Final"])
    rng = lambda col: f"'All 70 runs'!${col}$2:${col}${last}"
    totals = [("Valid", f"=COUNTIF({rng('G')},\"yes\")", f"=COUNTIF({rng('H')},\"yes\")"),
              ("All KPIs improved", f"=COUNTIF({rng('M')},\"yes\")", f"=COUNTIF({rng('N')},\"yes\")"),
              ("Valid and all KPIs improved", f"=COUNTIFS({rng('G')},\"yes\",{rng('M')},\"yes\")",
               f"=COUNTIFS({rng('H')},\"yes\",{rng('N')},\"yes\")"),
              ("Loop discarded a better initial", "", f"=COUNTIF({rng('U')},\"yes\")")]
    for k, (label, a, c) in enumerate(totals, t0 + 2):
        for j, v in enumerate((label, a, c), 1):
            ws1.cell(row=k, column=j, value=v if v != "" else "–").border = BOX
    ws1.cell(row=t0 + 6, column=1, value="Formulas count the rows of sheet \"All 70 runs\" (Excel recalculates them on opening).")
    k0 = t0 + 8
    ws1.cell(row=k0, column=1, value="Improved per KPI type").font = TITLE
    header(ws1, k0 + 1, ["KPI type", "Improved (initial)", "Improved (final)"])
    for k, t in enumerate(("Cycle time", "Waiting time", "Utilization"), k0 + 2):
        c = type_counts[t]
        for j, v in enumerate((t, f"{c['initial'][0]} of {c['initial'][1]}", f"{c['final'][0]} of {c['final'][1]}"), 1):
            ws1.cell(row=k, column=j, value=v).border = BOX
    ws1.cell(row=k0 + 5, column=1, value="Counts of KPI evaluations (10-seed) whose verdict is improved, from the run "
             "records. Time KPIs: GLASS's direction rule; utilization: the batch's rule (below 0.85 and lower than "
             "the baseline). A KPI is counted as waiting time when its name mentions waiting.")
    style_sheet(ws1, [22, 11, 13, 15, 12, 20, 11, 15, 15, 13, 11])

    # ------------------------------------------------------------- sheet 3
    ws3 = wb.create_sheet("Settings and scoring")
    header(ws3, 1, ["Item", "Value"])
    s = config["settings"]
    readme = (R / "README.md").read_text(encoding="utf-8")
    deviations = readme.split("## How GLASS was driven (differences from the UI)", 1)[1].split("\n## ", 1)[0].strip()
    rules = b.__doc__.split("SCORING", 1)[1].split("END SCORING")[0].strip("\n")
    items = [("Model", s["model"]),
             ("Temperature schedule", "0.3 for the initial proposal; GLASS raises it by 0.1 per iteration "
              "(min(0.3 + 0.1 × (iteration − 1), 0.7)): 0.3 / 0.4 / 0.5 / 0.6. GLASS's context-summary and "
              "compliance calls run at 0.1."),
             ("Maximum iterations", s["max_iterations"]), ("Repetitions per configuration", s["repetitions"]),
             ("Seeds", f"{s['seeds']} (42–51), paired: baseline, initial and final use the same seeds"),
             ("Cases per simulation", s["cases"]), ("Simulation start", s["start"]),
             ("GLASS commit of the runs", "8039058")]
    for proc, bc in config["business_context"].items():
        for key in ("staffing", "budget", "immutable"):
            items.append((f"Business context, {proc}: {key} (verbatim)", bc[key]))
        items.append((f"Hourly costs, {proc}", ", ".join(f"{k} {v} EUR" for k, v in bc["cost_per_hour"].items())))
    items += [("KPI validation (lead author's interpretation)",
               "Utilization KPIs were passed with target_direction minimize (Stage 1: maximize, maintain in one case) "
               "and the sentence \"" + b.OVERLOAD_TEXT + "\" appended to the description. Rationale: in the G3 "
               "configurations the target pool is overloaded; maximizing its utilization would reward overloading it "
               "and score the recovery (more staff) as a degradation."),
              ("Scoring rules (primary and secondary)", rules),
              ("Midnight mapping (distance measures only)",
               "Prosimos logs a task that ends exactly at the end of a working interval (e.g. 17:00) at 00:00 the next "
               "day. For the distance measures only, every log (baseline, initial, final, healthy) maps an end of "
               "exactly 00:00:00 that directly follows the resource's working-day end on the previous day back to that "
               "end; counts are recorded per log and the raw logs are unchanged. CWD on the raw logs is also recorded.")]
    items += [(f"GLASS fix {h}", t) for h, t in GLASS_FIXES]
    items.append(("Deviations from the GLASS UI (README)", deviations))
    for i, (k, v) in enumerate(items, 2):
        ws3.cell(row=i, column=1, value=k).border = BOX
        ws3.cell(row=i, column=2, value=v).border = BOX
    style_sheet(ws3, [38, 120])

    # ------------------------------------------------------------- sheet 4
    ws4 = wb.create_sheet("Findings")
    header(ws4, 1, ["#", "Finding", "Evidence"])
    for i, (f_, e_) in enumerate(FINDINGS, 2):
        for j, v in enumerate((i - 1, f_, e_), 1):
            ws4.cell(row=i, column=j, value=v).border = BOX
    style_sheet(ws4, [4, 50, 110])

    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    print(f"written {OUT.relative_to(P.REPO).as_posix()}; runs {last - 1}; per KPI type {type_counts}")


if __name__ == "__main__":
    main()

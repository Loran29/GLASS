"""Compare the manual Prosimos models with GLASS v3.2 (V_full, V_no_kb) and the baseline.

One KPI value per log (seeds 42-51), computed with GLASS's kpi_computation.py. Results go to results/. Reps where GLASS
kept the baseline have no final log, so the baseline values are used for them.

python -I scripts/compar_results_with_glass.py [--workers N] [--recompute] [--update-workbook] [--rep-sheets]
"""

from __future__ import annotations

import argparse
import glob
import importlib.util
import json
import math
import os
import re
import sys
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from scipy import stats

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "shared" / "loanapp_p2p"))
import lp_paths as P  # noqa: E402  (experiment paths)

COMP = Path(__file__).resolve().parents[1]
REPO = P.REPO
INPUTS = COMP / "inputs"
MINE = INPUTS / "manual_models"
V32_LOGS = INPUTS / "glass_event_logs" / "v3.2_rag_comparison"
RUNS = P.EXP2 / "results" / "stage2_results_v32_rag"
BASELINES = P.SHARED / "stage2" / "baselines"
KPI_MODULE = REPO / "goal_to_parameters" / "second_llm" / "kpi_computation.py"
OUT = COMP / "results"

# GLASS v3.2 with the knowledge base on (V_full) and off (V_no_kb).
ARMS = [
    dict(key="V_full", label="V_full v3.2", source="glass_v32_full", v32=True, suffix="_V_full_v3.2",
         sheet="Summary V_full v3.2", final_root=V32_LOGS / "V_full_system_simulated_eventlogs",
         records=RUNS / "V_full" / "raw"),
    dict(key="V_no_kb", label="V_no_kb v3.2", source="glass_v32_no_kb", v32=True, suffix="_V_no_kb_v3.2",
         sheet="Summary V_no_kb v3.2", final_root=V32_LOGS / "V_no_kb_simulated_eventlogs",
         records=RUNS / "V_no_kb" / "raw"),
]
SEEDS = list(range(42, 52))
SIM_START = "2023-02-09T10:00:00+02:00"
UTIL_LIMIT = 0.85
ALL_CONFIGS = [f"{p}_{g}" for p in ("LoanApp", "Procure2Pay")
               for g in ("G1", "G2", "G3", "G1G2", "G1G3", "G2G3", "G1G2G3")]

# Expected KPI per goal tag (EXPECTED_KPIS.md)
WAIT_ACTIVITY = {"LoanApp": "Assess loan risk", "Procure2Pay": "Release Supplier's Invoice"}
UTIL_ROLE = {"LoanApp": "Loan Officer", "Procure2Pay": "Purchasing Agent"}
CYCLE = "Average Cycle Time"


def expected_kpis(config: str) -> list[tuple[str, str, str]]:
    """[(kpi label, kind, target)] for a config, from its goal tags."""
    process, tags = config.split("_", 1)[0], re.findall(r"G\d", config.split("_", 1)[1])
    out = []
    if "G1" in tags:
        out.append((CYCLE, "cycle", ""))
    if "G2" in tags:
        out.append((f"{WAIT_ACTIVITY[process]} Waiting Time", "wait", WAIT_ACTIVITY[process]))
    if "G3" in tags:
        out.append((f"{UTIL_ROLE[process]} Utilization", "util", UTIL_ROLE[process]))
    return out


def norm(text: str) -> str:
    return text.replace("’", "'").replace("‘", "'").strip().lower()


def load_kpi_module():
    spec = importlib.util.spec_from_file_location("glass_kpi_computation", KPI_MODULE)
    mod = importlib.util.module_from_spec(spec)
    sys.modules["glass_kpi_computation"] = mod
    spec.loader.exec_module(mod)
    return mod


def compute_one(job: dict) -> list[dict]:
    km = load_kpi_module()
    df = pd.read_csv(job["log_path"])
    df = km._ensure_timestamps(df)
    rows = []
    base = {k: job[k] for k in ("source", "config", "variant", "model", "rep", "seed")}
    for label, kind, target in job["kpis"]:
        value, note = float("nan"), ""
        if kind == "cycle":
            k = km._compute_cycle_time(df)
            value = k.value if k else value
        elif kind == "wait":
            hit = [k for k in km._compute_activity_waiting_times(df) if norm(k.name) == norm(label)]
            if hit:
                value = hit[0].value
            else:
                note = "activity not found in log"
        elif kind == "util":
            k = km._compute_resource_utilization(df.copy(), job["params"], SIM_START)
            if k is None:
                note = "utilization not computed"
            else:
                per_role = k.details.get("per_role", {})
                roles = {norm(r): r for r in per_role}
                role = roles.get(norm(target))
                fb = k.details.get("fallback_span_based", [])
                if role is None:
                    note = f"role not found; roles={sorted(per_role)}"
                else:
                    members = [r for r, ro in k.details["resource_role"].items() if ro == role]
                    vals = [k.details["per_resource"][r] for r in members]
                    value = float(np.mean(vals))          # unrounded role mean
                    fb_role = [r for r in fb if r in members]
                    note = f"fallback_span_based={fb_role}" if fb_role else ""
                    if set(fb) - set(members):
                        note += f" other-role fallback={sorted(set(fb) - set(members))}"
        rows.append({**base, "kpi": label, "value": value, "note": note.strip()})
    return rows


def collect_jobs() -> tuple[list[dict], dict, dict]:
    """Logs to process, run info, and the reps per (arm, config) where GLASS kept the baseline (no final log)."""
    jobs: list[dict] = []
    info = {"mine_models": [], "skipped_configs": []}
    baseline_reps: dict[tuple[str, str], list[int]] = {}

    mine_models = []
    for d in sorted(glob.glob(str(MINE / "*" / "*"))):
        d = Path(d)
        if not d.is_dir() or d.parent.name.startswith("_") or d.parent.name == "comparison":
            continue
        js = list(d.glob("*.json"))
        if len(js) != 1:
            raise SystemExit(f"{d}: expected exactly one JSON, found {len(js)}")
        mine_models.append((d.parent.name, d.name, d.name.split("__")[0], json.loads(js[0].read_text(encoding="utf-8"))))
    info["mine_models"] = [(v, m, c) for v, m, c, _ in mine_models]
    configs = sorted({c for _, _, c, _ in mine_models}, key=ALL_CONFIGS.index)
    info["skipped_configs"] = [c for c in ALL_CONFIGS if c not in configs]

    for variant, model, config, params in mine_models:
        for s in SEEDS:
            p = MINE / variant / model / f"seed_{s}_log.csv"
            jobs.append(dict(source="mine", config=config, variant=variant, model=model, rep=0, seed=s,
                             log_path=str(p), params=params, kpis=expected_kpis(config)))

    for config in configs:
        bparams = json.loads((BASELINES / config / "sim_params.json").read_text(encoding="utf-8"))
        for s in SEEDS:
            jobs.append(dict(source="baseline", config=config, variant="", model="", rep=0, seed=s,
                             log_path=str(V32_LOGS / "baseline_logs" / config / f"seed_{s}.csv.gz"),
                             params=bparams, kpis=expected_kpis(config)))
        for arm in ARMS:
            baseline_reps[(arm["key"], config)] = []
            for rep in range(1, 6):
                rec = json.loads((arm["records"] / config / f"rep_{rep}.json").read_text(encoding="utf-8"))
                if arm["v32"]:
                    if rec["loop"].get("no_improving_scenario"):
                        baseline_reps[(arm["key"], config)].append(rep)
                        continue
                    final_iter = rec["evaluation"]["final"]["iteration"]
                else:
                    final_iter = rec["loop"]["best_iteration"]
                gparams = rec["loop"]["iterations"][final_iter - 1]["simulated_scenario_prosimos"]
                for s in SEEDS:
                    p = arm["final_root"] / config / f"rep_{rep}" / f"iter{final_iter}_seed{s}.csv.gz"
                    if not p.exists():
                        raise SystemExit(f"missing GLASS log {p} (final iteration in record = {final_iter})")
                    jobs.append(dict(source=arm["source"], config=config, variant="", model="", rep=rep, seed=s,
                                     log_path=str(p), params=gparams, kpis=expected_kpis(config)))
    return jobs, info, baseline_reps


def compare(a: np.ndarray, b: np.ndarray) -> dict:
    """a - b, paired by seed."""
    d = a - b
    n = len(d)
    mean_d = float(d.mean())
    sd = float(d.std(ddof=1)) if n > 1 else float("nan")
    se = sd / math.sqrt(n) if n > 1 else float("nan")
    tcrit = float(stats.t.ppf(0.975, n - 1)) if n > 1 else float("nan")
    lo, hi = mean_d - tcrit * se, mean_d + tcrit * se
    if n > 1 and sd > 0:
        p_paired = float(stats.ttest_rel(a, b).pvalue)
    else:
        p_paired = 0.0 if (n > 1 and mean_d != 0) else 1.0
    va, vb = a.std(ddof=1), b.std(ddof=1)
    p_welch = float(stats.ttest_ind(a, b, equal_var=False).pvalue) if (va > 0 or vb > 0) else (0.0 if mean_d != 0 else 1.0)
    return dict(a_mean=float(a.mean()), b_mean=float(b.mean()), mean_diff=mean_d, ci_low=float(lo), ci_high=float(hi),
                p_paired=p_paired, p_welch=p_welch)


def verdict(r: dict, a_name: str, b_name: str) -> str:
    if r["ci_low"] > 0:
        return f"{b_name} lower"
    if r["ci_high"] < 0:
        return f"{a_name} lower"
    return "no significant difference"


def pct_improvement(x: float, base: float) -> float:
    return float("nan") if base == 0 or math.isnan(base) else (base - x) / base * 100.0


EXCEL_COLS = ["config", "variant", "model", "kpi_expected", "baseline_mean", "mine_mean", "glass_mean_5reps",
              "glass_rep_min", "glass_rep_max", "mine_vs_baseline_pct", "glass_vs_baseline_pct",
              "diff_mine_minus_glass", "ci_low", "ci_high", "verdict"]
GLASS_COLS = EXCEL_COLS[6:]
BASELINE_FILL = "FFFCE4D6"               # light orange


def curated_layout(path: Path) -> tuple[list[tuple], dict]:
    """Row order and column widths of the first Summary sheet, so new sheets match it."""
    if not path.exists():
        return [], {}
    try:
        wb = load_workbook(path)
        name = next(n for n in wb.sheetnames if n.startswith("Summary"))
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        hdr = list(rows[0])
        ci, mi, ki = hdr.index("config"), hdr.index("model"), hdr.index("kpi_expected")
        order = [(r[ci], r[mi], r[ki]) for r in rows[1:] if r[ci]]
        widths = {k: v.width for k, v in ws.column_dimensions.items()}
        return order, widths
    except Exception:
        return [], {}


REP_COLS = (["config", "variant", "model", "kpi_expected", "baseline_mean", "manual_mean"]
            + [f"rep{k}_mean" for k in range(1, 6)]
            + ["reps_glass_better", "reps_no_significant_difference", "reps_manual_better"])


def utilization_counts(base: float, manual: float, reps: list[float]) -> tuple[int, int]:
    """(GLASS better, manual better) for utilization, from the means.

    Both below the baseline: GLASS is better if the rep is above the manual value, else manual is better.
    Only one below the baseline: that one is better. Neither: tie.
    """
    better = worse = 0
    m_ok = manual < base
    for g in reps:
        g_ok = g < base
        if g_ok and m_ok:
            better, worse = (better + 1, worse) if g > manual else (better, worse + 1)
        elif g_ok:
            better += 1
        elif m_ok:
            worse += 1
    return better, worse


def rep_sheet_df(arm: dict, all_df: pd.DataFrame, rank: dict) -> pd.DataFrame:
    """Per config x variant x KPI: manual mean, the 5 rep means, and how many reps are better, equal or worse.

    Time KPIs use the CI verdict; utilization uses utilization_counts.
    """
    d = all_df[(all_df.comparison == "mine vs GLASS") & all_df.b.str.match(r"GLASS rep\d$")]
    base = all_df[all_df.comparison == "mine vs baseline"].set_index(["config", "variant", "model", "kpi"]).b_mean
    rows = []
    for (config, variant, model, kpi), g in d.groupby(["config", "variant", "model", "kpi"], sort=False):
        r = dict(config=config, variant=variant, model=model, kpi_expected=kpi,
                 baseline_mean=float(base[(config, variant, model, kpi)]), manual_mean=float(g.a_mean.iloc[0]))
        for k in range(1, 6):
            r[f"rep{k}_mean"] = float(g[g.b == f"GLASS rep{k}"].b_mean.iloc[0])
        if kpi.endswith("Utilization"):
            better, manual = utilization_counts(r["baseline_mean"], r["manual_mean"], [r[f"rep{k}_mean"] for k in range(1, 6)])
        else:
            better = int(g.verdict.str.startswith("GLASS").sum())
            manual = int(g.verdict.str.startswith("mine").sum())
        r.update(reps_glass_better=better, reps_manual_better=manual, reps_no_significant_difference=5 - better - manual)
        rows.append(r)
    df = pd.DataFrame(rows, columns=REP_COLS)
    df["_o"] = [rank.get((r.config, r.model, r.kpi_expected), len(rank)) for r in df.itertuples()]
    return df.sort_values("_o", kind="stable").drop(columns="_o")


def fill_rep_sheet(ws, df: pd.DataFrame, arm: dict, baseline_reps: dict, widths: dict) -> None:
    ws.append(REP_COLS)
    for _, r in df.iterrows():
        ws.append([float(v) if isinstance(v, (np.floating, float)) else (int(v) if isinstance(v, np.integer) else v)
                   for v in (r[c] for c in REP_COLS)])
        for k in baseline_reps.get((arm["key"], r["config"]), []):
            ws.cell(row=ws.max_row, column=REP_COLS.index(f"rep{k}_mean") + 1).fill = PatternFill("solid", fgColor=BASELINE_FILL)
    n = len(df) * 5
    tot = [int(df[c].sum()) for c in ("reps_glass_better", "reps_no_significant_difference", "reps_manual_better")]
    ws.append([])
    ws.append(["TOTAL"] + [None] * (len(REP_COLS) - 4) + tot)
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append([f"{tot[0]} of {n} rep comparisons are better for GLASS, {tot[1]} not different, {tot[2]} better for manual."])
    ws.append(["Cycle time and waiting time: better means a lower mean with the 95% CI of the per-seed difference excluding zero."])
    ws.append(["Utilization: GLASS is better when manual < rep < baseline, or when only the rep is below the baseline; manual is "
               "better when the rep is at or below the manual mean, or when only the manual mean is below the baseline."])
    notes = baseline_footnotes(arm, baseline_reps)
    if notes:
        ws.append(["Orange: GLASS kept the baseline in that rep; the baseline value is used."])
        for line in notes:
            ws.append([line])
    for k, w in widths.items():
        if k in "ABCD":
            ws.column_dimensions[k].width = w
    for col in "EFGHIJKL":
        ws.column_dimensions[col].width = 18
    for col in "MNO":
        ws.column_dimensions[col].width = 30


def baseline_footnotes(arm: dict, baseline_reps: dict) -> list[str]:
    lines = []
    for (key, config), reps in baseline_reps.items():
        if key == arm["key"] and reps:
            lines.append(f"{config}: GLASS kept the baseline in rep{'s' if len(reps) > 1 else ''} "
                         f"{', '.join(map(str, reps))}; baseline values used.")
    return lines


def fill_sheet(ws, df: pd.DataFrame, flagged: set, notes: list[str], widths: dict) -> None:
    ws.append(EXCEL_COLS)
    for _, r in df.iterrows():
        ws.append([None if (isinstance(v, float) and math.isnan(v)) else (float(v) if isinstance(v, (np.floating,)) else v)
                   for v in (r[c] for c in EXCEL_COLS)])
        if (r["config"], r["model"], r["kpi_expected"]) in flagged:
            for col in range(EXCEL_COLS.index(GLASS_COLS[0]) + 1, len(EXCEL_COLS) + 1):
                ws.cell(row=ws.max_row, column=col).fill = PatternFill("solid", fgColor=BASELINE_FILL)
    if notes:
        ws.append([])
        ws.cell(row=ws.max_row + 1, column=1, value="Orange: baseline values used for the reps where GLASS kept the baseline.")
        for line in notes:
            ws.cell(row=ws.max_row + 1, column=1, value=line)
    for k, w in widths.items():
        ws.column_dimensions[k].width = w


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--workers", type=int, default=max(1, (os.cpu_count() or 2) - 2))
    ap.add_argument("--recompute", action="store_true",
                    help="recompute every KPI, ignoring the cache")
    ap.add_argument("--workbook", type=Path, default=OUT / "comparison_new.xlsx",
                    help="workbook for --update-workbook and --rep-sheets")
    ap.add_argument("--rep-sheets", action="store_true",
                    help="add the 'Reps vs manual' sheets to the workbook")
    ap.add_argument("--update-workbook", action="store_true",
                    help="add or replace the v3.2 summary sheets in the workbook")
    args = ap.parse_args()

    jobs, info, baseline_reps = collect_jobs()
    OUT.mkdir(parents=True, exist_ok=True)
    cache = OUT / "kpi_values_per_log.csv"
    cached = pd.DataFrame()
    if cache.exists() and not args.recompute:
        cached = pd.read_csv(cache).fillna({"variant": "", "model": "", "note": ""})
        cached["source"] = cached["source"].replace({"glass": "glass_v3.1"})
    key = lambda s, c, v, m, r, d: (s, c, v, m, int(r), int(d))
    done = {key(*t) for t in zip(cached.source, cached.config, cached.variant, cached.model, cached.rep, cached.seed)} \
        if len(cached) else set()
    todo = [j for j in jobs if key(j["source"], j["config"], j["variant"], j["model"], j["rep"], j["seed"]) not in done]
    print(f"{len(jobs)} logs needed, {len(jobs) - len(todo)} KPI sets cached, {len(todo)} to compute "
          f"with {args.workers} workers", flush=True)
    results: list[dict] = []
    if todo:
        with ProcessPoolExecutor(max_workers=args.workers) as ex:
            for i, rows in enumerate(ex.map(compute_one, todo, chunksize=4), 1):
                results.extend(rows)
                if i % 100 == 0:
                    print(f"  {i}/{len(todo)}", flush=True)
    vals = pd.concat([cached, pd.DataFrame(results)], ignore_index=True)
    vals.to_csv(cache, index=False)

    problems = vals[(vals["value"].isna()) | (vals["note"] != "")]
    bad = problems[problems["note"].str.contains("fallback_span_based|not found|not computed")
                   | problems["value"].isna()]

    def series(source, config, kpi, variant="", model="", rep=0):
        q = vals[(vals.source == source) & (vals.config == config) & (vals.kpi == kpi) & (vals.rep == rep)
                 & (vals.variant == variant) & (vals.model == model)].sort_values("seed")
        return q["value"].to_numpy(dtype=float)

    configs = sorted(vals.config.unique(), key=ALL_CONFIGS.index)
    summaries: dict[str, pd.DataFrame] = {}
    alls: dict[str, pd.DataFrame] = {}

    for arm in ARMS:
        all_rows, summary_rows = [], []

        def add(config, variant, model, kpi, comparison, a_name, b_name, a, b, extra=None):
            r = compare(a, b)
            row = dict(glass_arm=arm["label"], config=config, variant=variant, model=model, kpi=kpi,
                       comparison=comparison, a=a_name, b=b_name, **r, verdict=verdict(r, a_name, b_name))
            if extra:
                row.update(extra)
            all_rows.append(row)
            return row

        for config in configs:
            my_models = [(v, m) for v, m, c in info["mine_models"] if c == config]
            for label, kind, _ in expected_kpis(config):
                base = series("baseline", config, label)
                rep_series = {}
                for k in range(1, 6):
                    rep_series[k] = base if k in baseline_reps[(arm["key"], config)] \
                        else series(arm["source"], config, label, rep=k)
                glass_mean = np.mean([rep_series[k] for k in range(1, 6)], axis=0)
                gl = {**{f"GLASS rep{k}": rep_series[k] for k in rep_series}, "GLASS mean of 5 reps": glass_mean}
                # GLASS vs baseline (once per config)
                for name, arr in gl.items():
                    add(config, "", "", label, "GLASS vs baseline", name, "baseline", arr, base,
                        dict(improvement_pct_a=pct_improvement(arr.mean(), base.mean())))
                for variant, model in my_models:
                    mine = series("mine", config, label, variant, model)
                    add(config, variant, model, label, "mine vs baseline", "mine", "baseline", mine, base,
                        dict(improvement_pct_a=pct_improvement(mine.mean(), base.mean())))
                    for name, arr in gl.items():
                        add(config, variant, model, label, "mine vs GLASS", "mine", name, mine, arr,
                            dict(improvement_pct_a=pct_improvement(mine.mean(), base.mean()),
                                 improvement_pct_b=pct_improvement(arr.mean(), base.mean())))
                    # summary row
                    main_cmp = next(r for r in all_rows if r["config"] == config and r["model"] == model
                                    and r["kpi"] == label and r["comparison"] == "mine vs GLASS"
                                    and r["b"] == "GLASS mean of 5 reps")
                    rep_means = [rep_series[k].mean() for k in range(1, 6)]
                    srow = dict(glass_arm=arm["label"], config=config, variant=variant, model=model, kpi=label,
                                baseline_mean=base.mean(), mine_mean=mine.mean(), glass_mean_5reps=glass_mean.mean(),
                                glass_rep_min=min(rep_means), glass_rep_max=max(rep_means),
                                mine_vs_baseline_pct=pct_improvement(mine.mean(), base.mean()),
                                glass_vs_baseline_pct=pct_improvement(glass_mean.mean(), base.mean()),
                                diff_mine_minus_glass=main_cmp["mean_diff"], ci_low=main_cmp["ci_low"],
                                ci_high=main_cmp["ci_high"], p_paired=main_cmp["p_paired"], p_welch=main_cmp["p_welch"],
                                verdict=main_cmp["verdict"])
                    if kind == "util":
                        srow.update(mine_target_met=bool(mine.mean() < UTIL_LIMIT and mine.mean() < base.mean()),
                                    glass_target_met=bool(glass_mean.mean() < UTIL_LIMIT and glass_mean.mean() < base.mean()),
                                    baseline_target_met=bool(base.mean() < UTIL_LIMIT))
                    summary_rows.append(srow)
        summaries[arm["key"]] = pd.DataFrame(summary_rows)
        alls[arm["key"]] = pd.DataFrame(all_rows)

    # cross-check against the baseline values in the V_full run records
    checks = []
    for config in configs:
        rec = json.loads((ARMS[0]["records"] / config / "rep_1.json").read_text(encoding="utf-8"))
        final = rec["evaluation"]["final"]
        mine_series = {lab: series("baseline", config, lab) for lab, _, _ in expected_kpis(config)}
        for kp in final.get("kpis", []):
            bv = np.array(kp.get("baseline_values") or [], dtype=float)
            if len(bv) != 10:
                continue
            match = [lab for lab, s in mine_series.items() if np.allclose(s, bv, rtol=0.002, atol=0.011)]
            checks.append(dict(config=config, glass_record_kpi=kp["kpi"], category=kp["category"],
                               matches_my_series=", ".join(match) if match else "no match",
                               max_abs_diff_to_closest=float(min(np.abs(s - bv).max() for s in mine_series.values()))))
        ur = final.get("utilization_rule") or {}
        per_seed = ur.get("baseline_per_seed")
        if per_seed and any(k == "util" for _, k, _ in expected_kpis(config)):
            lab = next(l for l, k, _ in expected_kpis(config) if k == "util")
            s = mine_series[lab]
            checks.append(dict(config=config, glass_record_kpi="utilization_rule.baseline_per_seed", category="utilization",
                               matches_my_series=lab if np.allclose(s, per_seed, atol=0.002) else "no match",
                               max_abs_diff_to_closest=float(np.abs(s - np.array(per_seed)).max())))
    checks = pd.DataFrame(checks)

    for arm in ARMS:
        summaries[arm["key"]].to_csv(OUT / f"comparison_summary{arm['suffix']}.csv", index=False)
        alls[arm["key"]].to_csv(OUT / f"comparison_all{arm['suffix']}.csv", index=False)

    order, widths = curated_layout(args.workbook)
    rank = {k: i for i, k in enumerate(order)}

    def sheet_inputs(arm):
        df = summaries[arm["key"]].rename(columns={"kpi": "kpi_expected"})
        df = df.assign(_o=[rank.get((r.config, r.model, r.kpi_expected), len(rank)) for r in df.itertuples()])
        df = df.sort_values("_o", kind="stable").drop(columns="_o")
        flagged = {(config, model, kpi) for (key_, config), reps in baseline_reps.items()
                   if key_ == arm["key"] and reps
                   for model, kpi in df[df.config == config][["model", "kpi_expected"]].itertuples(index=False)}
        return df, flagged, baseline_footnotes(arm, baseline_reps)

    gen = Workbook()
    gen.remove(gen.active)
    for arm in ARMS:
        df, flagged, lines = sheet_inputs(arm)
        fill_sheet(gen.create_sheet(arm["sheet"]), df, flagged, lines, widths)
    gen.save(OUT / "comparison_generated.xlsx")
    if args.update_workbook:
        wb = load_workbook(args.workbook)
        for arm in ARMS:
            if arm["sheet"] in wb.sheetnames:
                del wb[arm["sheet"]]
            df, flagged, lines = sheet_inputs(arm)
            fill_sheet(wb.create_sheet(arm["sheet"]), df, flagged, lines, widths)
        wb.save(args.workbook)
        print("curated workbook updated:", wb.sheetnames)
    if args.rep_sheets:
        wb = load_workbook(args.workbook)
        for arm in ARMS:
            name = f"Reps vs manual {arm['label']}"
            if name in wb.sheetnames:
                del wb[name]
            fill_rep_sheet(wb.create_sheet(name), rep_sheet_df(arm, alls[arm["key"]], rank), arm, baseline_reps, widths)
        wb.save(args.workbook)
        print("rep sheets added:", wb.sheetnames)
    print("written to", OUT)
    if info["skipped_configs"]:
        print("not compared (no manual model):", ", ".join(info["skipped_configs"]))
    print(f"problems in KPI computation: {len(bad)}; cross-check rows without match: "
          f"{int((checks['matches_my_series'] == 'no match').sum()) if len(checks) else 0}")


if __name__ == "__main__":
    main()

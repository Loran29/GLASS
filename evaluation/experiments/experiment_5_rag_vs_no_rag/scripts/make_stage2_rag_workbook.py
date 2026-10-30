"""
Build results/GLASS_Stage2_RAG_comparison.xlsx from the verified files of the RAG comparison on GLASS v3.2
(values, no formulas).

Sources: results/stage2_rag_comparison.csv and .md, results/stage2_results_v32_rag/{V_full,V_no_kb}/ (records and
summary), results/VERSIONS.md (version table, copied verbatim), setup/stage2_config.json.
Sheets: RAG comparison per configuration (V_full v3.2 / V_no_kb v3.2, groups, tests), All 70 runs
(V_full v3.2), All 70 runs (V_no_kb v3.2), Settings and scoring.

Usage (repo root, environment with GLASS's dependencies and openpyxl):
  python evaluation/experiments/experiment_5_rag_vs_no_rag/scripts/make_stage2_rag_workbook.py
"""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "shared" / "loanapp_p2p"))
import lp_paths as P  # noqa: E402  (experiment paths; puts the other experiments' scripts on sys.path)
import make_stage2_workbook as w1  # noqa: E402  (names_for, readable, header, style_sheet, BOX, BOLD, TITLE)
import make_stage2_v1_v2_v3_v31_workbook as w31  # noqa: E402  (fmt, put, versions_table)
import run_stage2_batch as b  # noqa: E402
from compare_stage2_rag import direction_rejections  # noqa: E402

RES = HERE.parent / "results"
ARMS = {"V_full": P.EXP2 / "results" / "stage2_results_v32_rag" / "V_full",
        "V_no_kb": P.EXP2 / "results" / "stage2_results_v32_rag" / "V_no_kb"}
OUT = RES / "GLASS_Stage2_RAG_comparison.xlsx"
SOURCES = ("V_full", "V_no_kb")
CMP_COLS = [  # (csv metric, header)
    ("final_valid", "Valid scenario (after refinement)"),
    ("final_all_kpis_improved", "All target KPIs improved, our rule (after refinement)"),
    ("final_valid_and_improved", "Valid and all KPIs improved, our rule (after refinement)"),
    ("final_util_target_achieved_glass", "Utilization target achieved, GLASS: at most 0.85 (after refinement)"),
    ("final_util_target_achieved_batch", "Utilization target achieved, our rule: below 0.85 (after refinement)"),
    ("final_mean_recall", "Share of injected changes recovered, mean (after refinement)"),
    ("final_median_improvement_ratio_CTD", "Closeness to the healthy model, median (after refinement)"),
    ("reps_direction_rejection", "Runs with a direction_inconsistent rejection"),
    ("direction_rejections_per_run", "direction_inconsistent rejections per run"),
    ("reps_loop_discarded_better_initial", "Runs where the refinement made the result worse than the first proposal"),
    ("reps_loop_kept_baseline", "Runs where the loop kept the baseline (no improving scenario)"),
]
MEANS = {"final_mean_recall", "final_median_improvement_ratio_CTD", "direction_rejections_per_run"}
TARGET = {"final_util_target_achieved_glass", "final_util_target_achieved_batch"}
RUN_COLS = ["Configuration", "Process", "Problems", "Rep", "Refinement iterations used", "Final iteration",
            "Loop kept the baseline", "Valid (first proposal)", "Valid (after refinement)",
            "Violations (after refinement)", "KPI change % (first proposal)", "KPI change % (after refinement)",
            "All KPIs improved (first proposal)", "All KPIs improved (after refinement)",
            "Utilization target achieved, GLASS / our rule (after refinement)",
            "Injected changes recovered (after refinement)", "Closeness to healthy (after refinement)",
            "direction_inconsistent rejections", "Refinement made it worse", "Changes in the scenario after refinement"]


def runs_sheet(wb, title: str, folder: Path, config: dict) -> None:
    ws = wb.create_sheet(title)
    w1.header(ws, 1, RUN_COLS)
    summ = list(csv.DictReader(open(folder / "stage2_summary.csv", encoding="utf-8")))
    cfgs = {c["config_id"]: c for c in config["configurations"]}
    order = sorted(cfgs, key=lambda c: (c.split("_")[0], len(c.split("_")[1]), c))
    yes = {"True": "yes", "False": "no"}
    row = 2
    for cid in order:
        names, gs = w1.names_for(cfgs[cid], config), cfgs[cid]["injected_problems"]
        for rep in range(1, 6):
            rec = json.loads((folder / "raw" / cid / f"rep_{rep}.json").read_text(encoding="utf-8"))
            ev = rec.get("evaluation") or {}
            ph = {p: ev.get(p) or {"error": "no scenario"} for p in ("initial", "final")}
            srow = {x["phase"]: x for x in summ if x["config_id"] == cid and x["rep"] == str(rep)}
            loop = rec.get("loop") or {}
            kept = bool(loop.get("no_improving_scenario"))

            def yn(e):
                return "no scenario" if "error" in e else ("yes" if e.get("valid") else "no")

            def allimp(e):
                if "error" in e:
                    return "no scenario"
                return "yes" if e.get("kpis") and all(k.get("improved") is True for k in e["kpis"]) else "no"

            def kpis(e):
                if "error" in e:
                    return "no scenario"
                return "\n".join(f"{k['kpi']}: {k['pct_change']:+.1f} % ({'improved' if k.get('improved') else 'not improved'})"
                                 if k.get("pct_change") is not None else f"{k['kpi']}: n/a" for k in e.get("kpis") or [])

            def num(e, f):
                v = f(e) if "error" not in e else None
                return round(v, 4) if isinstance(v, (int, float)) else "–"

            s = srow.get("final", {})
            g, t = s.get("util_target_achieved_glass", ""), s.get("util_target_achieved_batch", "")
            target = "no scenario" if "error" in ph["final"] else ("–" if g == "" else f"{yes.get(g, g)} / {yes.get(t, t)}")
            it_no = ph["final"].get("iteration")
            fin_it = loop["iterations"][it_no - 1] if it_no else {}
            mods = "\n".join(
                f"{m['parameter_type']} {m['direction']}: {w1.readable(m['target_element'], names)}: "
                f"{w1.readable(str(m['baseline_value']), names)} → {w1.readable(str(m['proposed_value']), names)}"
                for m in ((fin_it or {}).get("patch") or {}).get("modifications", [])) or "–"
            w31.put(ws, row, [cid, rec["process"], "+".join(gs), rep, len(loop.get("iterations") or []), it_no or "–",
                              "yes" if kept else "no", yn(ph["initial"]), yn(ph["final"]),
                              "–" if "error" in ph["final"] else ("; ".join(ph["final"].get("violations") or []) or "none"),
                              kpis(ph["initial"]), kpis(ph["final"]), allimp(ph["initial"]), allimp(ph["final"]),
                              target, num(ph["final"], lambda e: (e.get("scores") or {}).get("recall")),
                              num(ph["final"], lambda e: (e.get("improvement_ratio") or {}).get("CTD")),
                              direction_rejections(rec),
                              "yes" if s.get("loop_discarded_better_initial") == "True" else "no", mods])
            row += 1
    w1.style_sheet(ws, [18, 12, 10, 5, 10, 9, 10, 10, 10, 26, 34, 34, 11, 11, 13, 10, 10, 11, 11, 60])


def main() -> None:
    b.quiet()
    config = json.loads(b.CONFIG_PATH.read_text(encoding="utf-8"))
    cmp_rows = list(csv.DictReader(open(RES / "stage2_rag_comparison.csv", encoding="utf-8")))
    md = (RES / "stage2_rag_comparison.md").read_text(encoding="utf-8")
    vmd = (P.EXP2 / "results" / "VERSIONS.md").read_text(encoding="utf-8")
    wb = Workbook()

    # ---------------------------------------------------------------- sheet 1: comparison
    ws = wb.active
    ws.title = "RAG comparison per configuration"
    w1.header(ws, 1, ["Configuration", "Has a utilization KPI"] + [h for _, h in CMP_COLS])
    r = 2
    for row in cmp_rows:
        util = row["has_utilization_kpi"]
        cells = [row["config_id"], {"True": "yes", "False": "no"}.get(util, "")]
        for m, _ in CMP_COLS:
            nd = 2 if m in MEANS else 0
            cells.append("–" if m in TARGET and util == "False"
                         else " / ".join(w31.fmt(row[f"{m}_{v}"], nd) for v in SOURCES))
        w31.put(ws, r, cells)
        if util == "":
            for j in range(1, len(cells) + 1):
                ws.cell(row=r, column=j).font = w1.BOLD
        r += 1
    notes = [line[2:].replace("**", "").replace("`", "") for line in md.split("\n\n| Configuration")[0].splitlines()
             if line.startswith("- ")]
    notes += [line[2:].replace("**", "") for line in
              md.split("**Utilization target achieved (GLASS, after refinement):**", 1)[1].split("\n\n")[0].splitlines()
              if line.startswith("- ")]
    notes.append("Source: stage2_rag_comparison.csv (from the verified per-configuration and summary CSVs and the "
                 "records of the two arms).")
    for n in notes:
        ws.cell(row=r, column=1, value=n)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="V_full vs V_no_kb (v3.2)").font = w1.TITLE
    test_part = md.split("**V_full vs V_no_kb (v3.2).**", 1)[1]
    ws.cell(row=r + 1, column=1, value=test_part.split("\n\n")[0].strip())
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=8)
    table = w31.versions_table(test_part)
    w1.header(ws, r + 2, table[0])
    for k, cells in enumerate(table[1:], r + 3):
        w31.put(ws, k, cells)
    r += 3 + len(table)
    for line in test_part.split("\n\n")[2].splitlines():
        if line.startswith("- "):
            ws.cell(row=r, column=1, value=line[2:].replace("**", ""))
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
            r += 1
    w1.style_sheet(ws, [22, 11] + [15] * len(CMP_COLS))

    # ---------------------------------------------------------------- sheets 2 and 3: all runs per arm
    for arm, folder in ARMS.items():
        runs_sheet(wb, f"All 70 runs ({arm} v3.2)", folder, config)

    # ---------------------------------------------------------------- sheet 4: settings
    ws3 = wb.create_sheet("Settings and scoring")
    w1.header(ws3, 1, ["Item", "Value"])
    s = config["settings"]
    rec = json.loads(next((ARMS["V_full"] / "raw").glob("*/rep_1.json")).read_text(encoding="utf-8"))
    rules = b.__doc__.split("SCORING", 1)[1].split("END SCORING")[0].strip()
    items = [
        ("How to read", "\"First proposal\" = the scenario GLASS proposes first; \"after refinement\" = the scenario "
         "GLASS selects after its simulation-based refinement (up to 4 iterations), or no scenario when the loop kept "
         "the baseline. On the first sheet each cell is V_full v3.2 / V_no_kb v3.2; counts are out of 5 runs "
         "per configuration."),
        ("Arms", "V_full: the Stage-2 batch on GLASS v3.2 with the knowledge base. V_no_kb: the same batch with the "
         "knowledge base off in generation and feedback (knowledge.retrieval.retrieve_for_second_llm returns an empty "
         "result, the evidence's kb_json is empty, the feedback's literature section is empty); the switch is the "
         "real-life ablation's (run_ablation_batch.py, disable_kb)."),
        ("GLASS commit (records)", f"{rec['glass_commit']} (goal_to_parameters/ identical to c2f7f53, v3.2)"),
        ("Batch script SHA-256", rec["script_sha256"]),
        ("Configuration SHA-256", rec["config_sha256"]),
        ("Model", s["model"]),
        ("Temperature schedule", "0.3 for the initial proposal; GLASS raises it by 0.1 per iteration "
         "(min(0.3 + 0.1 × (iteration − 1), 0.7)). GLASS's context-summary and compliance calls run at 0.1."),
        ("Maximum iterations", s["max_iterations"]), ("Repetitions per configuration", s["repetitions"]),
        ("Seeds", f"{s['seeds']} (42–51), paired: baseline, first proposal and after refinement use the same seeds"),
        ("Cases per simulation", s["cases"]), ("Simulation start", s["start"]),
        ("Loop kept the baseline (v3.2)", "When no iteration reaches a positive score, GLASS keeps the baseline: no "
         "final scenario, scored not valid. Such a run counts as \"refinement made it worse\" only if the first "
         "proposal was valid and improved at least one KPI on 10 seeds."),
        ("direction_inconsistent (v3.2)", "A modification whose stated direction contradicts its values is rejected "
         "and fed back; on the last attempt only that modification is dropped."),
        ("Scoring rules (primary and secondary)", rules),
    ]
    for i, (k, v) in enumerate(items, 2):
        w31.put(ws3, i, [k, v])
    r = len(items) + 3
    ws3.cell(row=r, column=1, value="Version table (VERSIONS.md, verbatim)").font = w1.TITLE
    table = w31.versions_table(vmd.split("**Rationale")[0])
    w1.header(ws3, r + 1, table[0])
    for k, cells in enumerate(table[1:], r + 2):
        w31.put(ws3, k, cells)
    w1.style_sheet(ws3, [38, 60, 60, 45, 45, 60, 22])

    wb.save(OUT)
    print(f"written {OUT.relative_to(P.REPO).as_posix()}")


if __name__ == "__main__":
    main()
